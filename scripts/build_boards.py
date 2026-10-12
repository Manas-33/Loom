#!/usr/bin/env python3
"""
Build and refresh the list of company job boards that Loom checks for new jobs.

Sources (all merged; each board remembers which sources listed it):
  aggregator   Feashliaa/job-board-aggregator company lists (MIT), incl. Workday
  findjobs     blakegrudzien/find-jobs company lists (names only)
  commoncrawl  board links in the last few Common Crawl snapshots
  simplify     apply links in the SimplifyJobs New-Grad-Positions list
  sheet        Title_URL / Apply_URL of every job in jobs.xlsx (pastes, sweeps)

New boards are checked once through the site's API and marked ok, dead,
closed (Workday site not taking applicants right now) or error. Later runs
recheck closed and error boards, and skip dead ones unless --recheck-dead.

Output: data/boards.json (gitignored), plus data/boards.csv listing every live
or closed board with its link, for browsing in a spreadsheet. Common Crawl's index server is often
slow or down, so it gets a time budget (--cc-minutes) and every downloaded
page is cached in data/cache/; old snapshots never change, and the next run
picks up where this one stopped.

Usage (from repo root):
  python scripts/build_boards.py                      # full refresh (monthly)
  python scripts/build_boards.py --sources sheet,simplify
  python scripts/build_boards.py --stats              # counts only
  python scripts/build_boards.py --export-only        # just rewrite data/boards.csv
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import tempfile
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime
from pathlib import Path

import requests
from urllib.parse import quote

sys.path.insert(0, str(Path(__file__).resolve().parent))
from boards import ATS, USER_AGENT, board_from_url, probe  # noqa: E402

REPO_ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = REPO_ROOT / "data"
BOARDS_FILE = DATA_DIR / "boards.json"
BOARDS_CSV = DATA_DIR / "boards.csv"
CC_CACHE = DATA_DIR / "cache" / "commoncrawl"

# Common Crawl goes last: its index server is slow and often down, so it runs on a time budget
SOURCES = ("aggregator", "findjobs", "simplify", "sheet", "commoncrawl")
AGGREGATOR = "https://raw.githubusercontent.com/Feashliaa/job-board-aggregator/main/data/{ats}_companies.json"
FINDJOBS = "https://raw.githubusercontent.com/blakegrudzien/find-jobs/main/companies_{ats}.txt"
SIMPLIFY = "https://raw.githubusercontent.com/SimplifyJobs/New-Grad-Positions/dev/README.md"
CC_INDEXES = "https://index.commoncrawl.org/collinfo.json"
CC_HOSTS = ("jobs.ashbyhq.com", "boards.greenhouse.io", "job-boards.greenhouse.io", "jobs.lever.co")

# Parallel checks per site; Ashby rate-limits hardest
WORKERS = {"greenhouse": 6, "lever": 12, "ashby": 4, "workday": 8}


def log(msg: str) -> None:
    print(f"[{datetime.now():%H:%M:%S}] {msg}", flush=True)


def _get(url: str, timeout: int = 60, tries: int = 3) -> requests.Response:
    last = None
    for attempt in range(tries):
        try:
            resp = requests.get(url, headers={"User-Agent": USER_AGENT}, timeout=timeout)
            if resp.status_code == 200:
                return resp
            last = f"HTTP {resp.status_code}"
        except requests.RequestException as exc:
            last = str(exc)
        time.sleep(5 * (attempt + 1))
    raise RuntimeError(f"{url}: {last}")


# ── Sources ──────────────────────────────────────────────────────────

def from_aggregator() -> dict[str, set[str]]:
    found: dict[str, set[str]] = {ats: set() for ats in ATS}
    for ats in ATS:
        names = _get(AGGREGATOR.format(ats=ats)).json()
        for name in names:
            name = str(name).strip().lower()
            if ats == "workday":
                if re.fullmatch(r"[a-z0-9-]+\|wd\d+\|[a-z0-9_-]+", name):
                    found[ats].add(name)
            elif re.fullmatch(r"[a-z0-9][a-z0-9 ._-]{0,80}", name):
                found[ats].add(name)
    return found


def from_findjobs() -> dict[str, set[str]]:
    found: dict[str, set[str]] = {ats: set() for ats in ATS}
    for ats in ("greenhouse", "ashby", "lever"):
        text = _get(FINDJOBS.format(ats=ats)).text
        for line in text.splitlines():
            name = line.split("#")[0].strip().lower()
            if name and re.fullmatch(r"[a-z0-9][a-z0-9 ._-]{0,80}", name):
                found[ats].add(name)
    return found


def from_links(links: list[str]) -> dict[str, set[str]]:
    found: dict[str, set[str]] = {ats: set() for ats in ATS}
    for link in links:
        board = board_from_url(link)
        if board:
            found[board[0]].add(board[1])
    return found


def from_simplify() -> dict[str, set[str]]:
    text = _get(SIMPLIFY).text
    return from_links(re.findall(r'href="(https?://[^"]+)"', text))


def from_sheet() -> dict[str, set[str]]:
    from openpyxl import load_workbook

    xlsx = REPO_ROOT / "jobs.xlsx"
    if not xlsx.exists():
        return {ats: set() for ats in ATS}
    ws = load_workbook(xlsx, read_only=True).active
    rows = ws.iter_rows(values_only=True)
    headers = [str(h or "") for h in next(rows)]
    cols = [headers.index(c) for c in ("Title_URL", "Apply_URL") if c in headers]
    links = [str(row[i]) for row in rows for i in cols if i < len(row) and row[i]]
    return from_links(links)


def from_commoncrawl(snapshots: int, minutes: float) -> dict[str, set[str]]:
    found: dict[str, set[str]] = {ats: set() for ats in ATS}
    deadline = time.monotonic() + minutes * 60
    crawls = [c["id"] for c in _get(CC_INDEXES, timeout=60, tries=2).json()[:snapshots]]
    CC_CACHE.mkdir(parents=True, exist_ok=True)
    for crawl in crawls:
        for host in CC_HOSTS:
            links, note = _cc_host(crawl, host, deadline)
            part = from_links(links)
            log(f"  Common Crawl {crawl} {host}: {sum(len(v) for v in part.values())} boards ({note})")
            for ats in ATS:
                found[ats] |= part[ats]
    return found


def _cc_host(crawl: str, host: str, deadline: float) -> tuple[list[str], str]:
    """Links for one host in one snapshot, one page at a time, cached per page."""
    legacy = CC_CACHE / f"{crawl}__{host}.txt"
    if legacy.exists():
        return legacy.read_text(encoding="utf-8").splitlines(), "cached"
    folder = CC_CACHE / f"{crawl}__{host}"
    folder.mkdir(exist_ok=True)
    if (folder / "complete").exists():
        return _cached_pages(folder), "cached"

    api = f"https://index.commoncrawl.org/{crawl}-index?url={host}/*&output=json&fl=url"
    pages_file = folder / "pages"
    if not pages_file.exists():
        if time.monotonic() > deadline:
            return [], "not started: out of time"
        try:
            pages_file.write_text(str(_get(f"{api}&showNumPages=true", timeout=90, tries=2).json()["pages"]))
        except (RuntimeError, ValueError, KeyError) as exc:
            return [], f"index unavailable: {str(exc)[-60:]}"
    pages = int(pages_file.read_text())

    for page in range(pages):
        cache = folder / f"page-{page}.txt"
        if cache.exists():
            continue
        if time.monotonic() > deadline:
            break
        try:
            resp = _get(f"{api}&page={page}", timeout=90, tries=2)
        except RuntimeError as exc:
            log(f"    page {page + 1}/{pages} failed ({str(exc)[-40:]}); stopping this host until next run")
            break
        boards: dict[tuple[str, str], str] = {}
        for line in resp.text.splitlines():
            try:
                board = board_from_url(json.loads(line)["url"])
            except (json.JSONDecodeError, KeyError):
                continue
            if board and board not in boards:
                boards[board] = json.loads(line)["url"]  # original link keeps Lever's case
        cache.write_text("\n".join(sorted(boards.values())), encoding="utf-8")

    have = sum(1 for p in range(pages) if (folder / f"page-{p}.txt").exists())
    if have == pages:
        (folder / "complete").touch()
        return _cached_pages(folder), "complete"
    return _cached_pages(folder), f"partial: {have}/{pages} pages, resumes next run"


def _cached_pages(folder: Path) -> list[str]:
    links: list[str] = []
    for page in sorted(folder.glob("page-*.txt")):
        links += [l for l in page.read_text(encoding="utf-8").splitlines() if l]
    return links


# ── The list file ────────────────────────────────────────────────────

def load_boards() -> dict:
    if BOARDS_FILE.exists():
        return json.loads(BOARDS_FILE.read_text(encoding="utf-8"))
    return {"version": 1, "updated": None, "boards": {ats: {} for ats in ATS}}


def save_boards(data: dict) -> None:
    DATA_DIR.mkdir(exist_ok=True)
    data["updated"] = datetime.now().isoformat(timespec="seconds")
    fd, tmp = tempfile.mkstemp(dir=DATA_DIR, prefix=".boards-")
    with os.fdopen(fd, "w", encoding="utf-8") as fh:
        json.dump(data, fh, indent=1, sort_keys=True)
    os.replace(tmp, BOARDS_FILE)


def merge(data: dict, source: str, found: dict[str, set[str]]) -> int:
    today = date.today().isoformat()
    new = 0
    for ats, keys in found.items():
        table = data["boards"].setdefault(ats, {})
        for key in keys:
            entry = table.get(key)
            if entry is None:
                table[key] = {"sources": [source], "added": today, "status": "new"}
                new += 1
            elif source not in entry["sources"]:
                entry["sources"].append(source)
    return new


def check_boards(data: dict, recheck_dead: bool, limit: int | None) -> None:
    todo = []
    for ats in ATS:
        for key, entry in data["boards"].get(ats, {}).items():
            if entry["status"] in ("new", "closed", "error") or (recheck_dead and entry["status"] == "dead"):
                todo.append((ats, key))
    if limit:
        todo = todo[:limit]
    if not todo:
        log("No boards to check.")
        return
    log(f"Checking {len(todo)} boards: " + ", ".join(f"{a} {sum(1 for t in todo if t[0] == a)}" for a in ATS))

    today = date.today().isoformat()
    pools = {ats: ThreadPoolExecutor(max_workers=WORKERS[ats]) for ats in ATS}
    futures = {pools[ats].submit(probe, ats, key): (ats, key) for ats, key in todo}
    done = 0
    try:
        for fut in as_completed(futures):
            ats, key = futures[fut]
            result = fut.result()
            entry = data["boards"][ats][key]
            entry.update(status=result.status, checked=today)
            if result.jobs is not None:
                entry["jobs"] = result.jobs
            if result.detail:
                entry["detail"] = result.detail
            else:
                entry.pop("detail", None)
            done += 1
            if done % 1000 == 0:
                save_boards(data)
                log(f"  {done}/{len(todo)} checked")
    finally:
        for pool in pools.values():
            pool.shutdown(wait=True, cancel_futures=True)
        save_boards(data)


def board_link(ats: str, key: str) -> str:
    if ats == "workday":
        tenant, wd, site = key.split("|")
        return f"https://{tenant}.{wd}.myworkdayjobs.com/{site}"
    base = {"greenhouse": "https://job-boards.greenhouse.io", "ashby": "https://jobs.ashbyhq.com", "lever": "https://jobs.lever.co"}[ats]
    return f"{base}/{quote(key)}"


# Board names that don't read as the company (mostly Workday tenant codes)
KNOWN_NAMES = {
    "globalhr": "RTX", "ngc": "Northrop Grumman", "caci": "CACI", "pae": "Amentum (PAE)", "hpe": "Hewlett Packard Enterprise",
    "bah": "Booz Allen Hamilton", "capitalone": "Capital One", "ghr": "Lockheed Martin", "pwc": "PwC",
    "wgu": "Western Governors University", "gevernova": "GE Vernova", "relx": "RELX", "thermofisher": "Thermo Fisher Scientific",
    "usaa": "USAA", "aes": "AES", "btisolutions": "BTI Solutions", "draper": "Draper", "jeffersonhealth": "Jefferson Health",
    "aero": "The Aerospace Corporation", "bnl": "Brookhaven National Laboratory", "fis": "FIS", "kbr": "KBR",
    "mufgub": "MUFG", "radiancetech": "Radiance Technologies", "umd": "University of Maryland", "aoins": "Auto-Owners Insurance",
    "nvidia": "NVIDIA", "gm": "General Motors", "ibm": "IBM", "jpmc": "JPMorgan Chase", "bofa": "Bank of America",
}


def readable_name(ats: str, key: str) -> str:
    """A readable company name from the board name; the sites don't give the real one."""
    raw = key.split("|")[0] if ats == "workday" else key
    if raw.lower() in KNOWN_NAMES:
        return KNOWN_NAMES[raw.lower()]
    if raw != raw.lower():  # Lever names keep the company's own capitalization
        return raw
    words = re.split(r"[-_ ]+", raw)
    return " ".join("AI" if w == "ai" else w.capitalize() for w in words if w)


def export_csv(data: dict) -> int:
    import csv

    rows = []
    for ats in ATS:
        for key, e in data["boards"].get(ats, {}).items():
            if e["status"] not in ("ok", "closed"):
                continue
            rows.append({
                "company": readable_name(ats, key),
                "site": ats,
                "board": key,
                "link": board_link(ats, key),
                "status": "live" if e["status"] == "ok" else "closed for now",
                "open_jobs": e.get("jobs", ""),
                "found_in": ", ".join(e["sources"]),
                "added": e.get("added", ""),
                "last_checked": e.get("checked", ""),
            })
    rows.sort(key=lambda r: (r["company"].lower(), r["site"]))
    DATA_DIR.mkdir(exist_ok=True)
    with open(BOARDS_CSV, "w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(rows[0]) if rows else ["company"])
        writer.writeheader()
        writer.writerows(rows)
    return len(rows)


def stats(data: dict) -> str:
    lines = []
    for ats in ATS:
        table = data["boards"].get(ats, {})
        by_status: dict[str, int] = {}
        for entry in table.values():
            by_status[entry["status"]] = by_status.get(entry["status"], 0) + 1
        lines.append(f"  {ats:10} {len(table):6} boards  " + "  ".join(f"{k} {v}" for k, v in sorted(by_status.items())))
    return "\n".join(lines)


def main() -> int:
    parser = argparse.ArgumentParser(description="Build the company job-board list.")
    parser.add_argument("--sources", default=",".join(SOURCES), help=f"comma list from: {', '.join(SOURCES)}")
    parser.add_argument("--snapshots", type=int, default=4, help="Common Crawl snapshots to read (default 4)")
    parser.add_argument("--cc-minutes", type=float, default=15, help="time budget for Common Crawl (default 15)")
    parser.add_argument("--no-check", action="store_true", help="merge sources without checking new boards")
    parser.add_argument("--recheck-dead", action="store_true", help="also recheck boards marked dead")
    parser.add_argument("--limit", type=int, default=None, help="check at most this many boards (testing)")
    parser.add_argument("--stats", action="store_true", help="print counts and exit")
    parser.add_argument("--export-only", action="store_true", help="rewrite data/boards.csv and exit")
    args = parser.parse_args()

    data = load_boards()
    if args.stats:
        print(f"updated {data.get('updated')}\n{stats(data)}")
        return 0
    if args.export_only:
        print(f"Wrote {export_csv(data)} boards to {BOARDS_CSV.relative_to(REPO_ROOT)}")
        return 0

    chosen = [s.strip() for s in args.sources.split(",") if s.strip()]
    unknown = set(chosen) - set(SOURCES)
    if unknown:
        parser.error(f"unknown source(s): {', '.join(sorted(unknown))}")

    readers = {
        "aggregator": from_aggregator,
        "findjobs": from_findjobs,
        "commoncrawl": lambda: from_commoncrawl(args.snapshots, args.cc_minutes),
        "simplify": from_simplify,
        "sheet": from_sheet,
    }
    for source in chosen:
        log(f"Reading {source}…")
        try:
            found = readers[source]()
        except (RuntimeError, ValueError, requests.RequestException) as exc:
            log(f"  {source} failed: {exc}")
            continue
        new = merge(data, source, found)
        save_boards(data)
        log(f"  {source}: {sum(len(v) for v in found.values())} boards, {new} new")

    if not args.no_check:
        check_boards(data, args.recheck_dead, args.limit)
    count = export_csv(data)
    log(f"Done. {count} live or closed boards written to {BOARDS_CSV.relative_to(REPO_ROOT)}\n" + stats(data))
    return 0


if __name__ == "__main__":
    sys.exit(main())
