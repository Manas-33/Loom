#!/usr/bin/env python3
"""
Single entry point for adding a job to Loom, whatever found it.

Every source (job-board checks, SimplifyJobs, LinkedIn/Handshake sweeps,
pasted text) hands a job record here. add_job() validates it, checks for a
duplicate (job ID from the apply/posting link first, then company + title),
and either merges the new sources into the existing row or writes a new
jobs.xlsx row, jobs/<slug>.md and the dashboard status.

Record format (JSON keys):
  company, title            required
  location                  "Not specified" when unknown
  url                       where the posting was found (LinkedIn, board, ...)
  apply_url                 the real application link on the company site
  description               full job description text
  found_via                 list: linkedin, handshake, greenhouse, lever,
                            ashby, workday, simplify, paste
  apply_via                 greenhouse, lever, ashby, workday, easy-apply,
                            handshake, other (derived from apply_url if unset)
  easy_apply                true when LinkedIn shows Easy Apply
  category, salary, posted_time, people_applied, company_url, company_size,
  company_followers, company_intro, internship (Yes/No/Maybe)
  early_career (Yes/No), years_required, grad_year, flags (list),
  suggested_bucket          filled by later pipeline steps

Usage (from repo root):
  python scripts/add_job.py record.json          # object, list, or JSONL
  cat records.jsonl | python scripts/add_job.py -
  python scripts/add_job.py record.json --dry-run
Prints one JSON result per record: added / merged / duplicate / rejected.
"""

from __future__ import annotations

import argparse
import contextlib
import fcntl
import json
import os
import re
import shutil
import sys
import tempfile
import uuid
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterator

REPO_ROOT = Path(__file__).resolve().parent.parent

FOUND_VIA = ("linkedin", "handshake", "greenhouse", "lever", "ashby", "workday", "simplify", "paste")
APPLY_VIA = ("greenhouse", "lever", "ashby", "workday", "easy-apply", "handshake", "other")
INTERNSHIP = ("Yes", "No", "Maybe")
DEFAULT_STATUS = "Not Applied"
UNKNOWN_LOCATION = "Not specified"

# record key -> jobs.xlsx column. The first 17 are the sheet's original columns;
# the rest are appended to the header row the first time add_job runs.
COLUMNS: dict[str, str] = {
    "company": "Company",
    "title": "Title",
    "category": "Category",
    "location": "Location",
    "url": "Title_URL",
    "status": "Status",
    "date": "Date",
    "description": "About_the_job",
    "posted_time": "Posted_time",
    "salary": "Salary",
    "people_applied": "People_applied",
    "company_url": "Company_URL",
    "company_followers": "Company_follower",
    "company_size": "Company_size",
    "linkedin_employees": "Count_of_employee_onLinkedIn",
    "company_intro": "Company_Intro",
    "internship": "Is this an internship",
    "found_via": "Found_via",
    "apply_url": "Apply_URL",
    "apply_via": "Apply_via",
    "job_id": "Job_ID",
    "early_career": "Early_career",
    "years_required": "Years_required",
    "grad_year": "Grad_year",
    "flags": "Flags",
    "suggested_bucket": "Suggested_bucket",
}
LIST_FIELDS = ("found_via", "flags")
# Accepted but not stored in the sheet: an input hint, and the inbox's new-grad / general title tag
NOT_STORED = ("easy_apply", "title_match")
LONG_FIELDS = ("description", "company_intro")
MAX_LEN = {"description": 60000, "company_intro": 8000}
DEFAULT_MAX_LEN = 500


class InvalidJob(ValueError):
    """The record can't be stored as given."""


# ── Job IDs and links ─────────────────────────────────────────────────

_ID_PATTERNS: list[tuple[str, re.Pattern[str]]] = [
    ("greenhouse", re.compile(r"[?&]gh_jid=(\d+)")),
    ("greenhouse", re.compile(r"greenhouse\.io/[^/?#]+/jobs/(\d+)")),
    ("greenhouse", re.compile(r"greenhouse\.io/embed/job_app\?.*?token=(\d+)")),
    ("ashby", re.compile(r"jobs\.ashbyhq\.com/[^/?#]+/([0-9a-f]{8}-[0-9a-f-]{27})", re.I)),
    ("lever", re.compile(r"jobs\.lever\.co/[^/?#]+/([0-9a-f]{8}-[0-9a-f-]{27})", re.I)),
    ("linkedin", re.compile(r"linkedin\.com/jobs/view/(?:[^/?#]*?-)?(\d{6,})")),
    ("linkedin", re.compile(r"linkedin\.com/jobs/.*?[?&]currentJobId=(\d{6,})")),
    ("handshake", re.compile(r"joinhandshake\.com/(?:stu/|job-search/)?jobs?/(\d+)")),
    ("handshake", re.compile(r"joinhandshake\.com/job-search/(\d+)")),
    ("indeed", re.compile(r"indeed\.com/.*?[?&]jk=([0-9a-f]+)", re.I)),
    ("smartrecruiters", re.compile(r"jobs\.smartrecruiters\.com/[^/?#]+/(\d+)")),
    ("greenhouse", re.compile(r"careerpuck\.com/job-board/[^/?#]+/job/(\d+)")),  # Greenhouse boards on Careerpuck
]
_WORKDAY = [
    re.compile(r"https?://([a-z0-9-]+)\.wd\d+\.myworkdayjobs\.com/(?:[^?#]*/)?job/[^?#]*?_([A-Za-z0-9-]+)/?(?:[?#]|$)", re.I),
    re.compile(r"https?://wd\d+\.myworkdaysite\.com/recruiting/([a-z0-9-]+)/[^?#]*?/job/[^?#]*?_([A-Za-z0-9-]+)/?(?:[?#]|$)", re.I),
]


def _canonical_id(job_id: str) -> str:
    """Workday reposts one opening as R123, R123-1, R123-2 (other sites or brands): they're one job."""
    return re.sub(r"(\d)-\d{1,2}$", r"\1", job_id) if job_id.startswith("workday:") else job_id


def job_id_from_url(url: str) -> str | None:
    """Stable ID for a posting link, e.g. 'ashby:<uuid>', or None."""
    if not url:
        return None
    for pattern in _WORKDAY:
        m = pattern.search(url)
        if m:
            return _canonical_id(f"workday:{m.group(1).lower()}:{m.group(2).lower()}")
    for site, pattern in _ID_PATTERNS:
        m = pattern.search(url)
        if m:
            return f"{site}:{m.group(1).lower()}"
    return None


def apply_via_from_url(url: str) -> str:
    host = re.sub(r"^https?://", "", url or "").split("/")[0].lower()
    if "greenhouse.io" in host or "gh_jid=" in (url or ""):
        return "greenhouse"
    if host.endswith("lever.co"):
        return "lever"
    if host.endswith("ashbyhq.com"):
        return "ashby"
    if "myworkdayjobs.com" in host:
        return "workday"
    if host.endswith("joinhandshake.com"):
        return "handshake"
    return "other"


# ── Normalizing and validating a record ──────────────────────────────

_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def _clean_text(value: Any, key: str) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value != value:  # NaN from pandas
        return ""
    text = _CONTROL.sub("", str(value)).strip()
    if key not in LONG_FIELDS:
        text = re.sub(r"\s+", " ", text)
    limit = MAX_LEN.get(key, DEFAULT_MAX_LEN)
    if len(text) > limit:
        raise InvalidJob(f"{key} is longer than {limit} characters")
    return text


def _clean_list(value: Any, key: str) -> list[str]:
    if value is None or value == "":
        return []
    items = value if isinstance(value, (list, tuple)) else re.split(r"[,;]", str(value))
    out: list[str] = []
    for item in items:
        token = _clean_text(item, key).lower()
        if not token:
            continue
        if not re.fullmatch(r"[a-z0-9][a-z0-9-]{0,40}", token):
            raise InvalidJob(f"{key} value {token!r} isn't a short lowercase tag")
        if token not in out:
            out.append(token)
    return out


def _check_url(url: str, key: str) -> str:
    if url and not re.match(r"https?://[^\s/]+\.[^\s/]+", url):
        raise InvalidJob(f"{key} must be an http(s) link, got {url[:80]!r}")
    return url


def normalize(record: dict[str, Any]) -> dict[str, Any]:
    """Return a clean copy of a record, or raise InvalidJob."""
    if not isinstance(record, dict):
        raise InvalidJob("record must be a JSON object")
    record = {k: v for k, v in record.items() if not str(k).startswith("_")}  # pipeline notes
    unknown = set(record) - set(COLUMNS) - set(NOT_STORED)
    if unknown:
        raise InvalidJob(f"unknown fields: {', '.join(sorted(unknown))}")

    job: dict[str, Any] = {}
    for key in COLUMNS:
        if key in LIST_FIELDS:
            job[key] = _clean_list(record.get(key), key)
        else:
            job[key] = _clean_text(record.get(key), key)

    if not job["company"]:
        raise InvalidJob("company is required")
    if not job["title"]:
        raise InvalidJob("title is required")
    if not job["location"] or job["location"].lower() in ("nan", "none", "n/a"):
        job["location"] = UNKNOWN_LOCATION

    for key in ("url", "apply_url", "company_url"):
        _check_url(job[key], key)

    bad = [v for v in job["found_via"] if v not in FOUND_VIA]
    if bad:
        raise InvalidJob(f"found_via has unknown source(s): {', '.join(bad)}")
    if not job["found_via"]:
        raise InvalidJob("found_via is required (where the job came from)")

    if job["internship"] and job["internship"] not in INTERNSHIP:
        raise InvalidJob(f"internship must be one of {', '.join(INTERNSHIP)}")
    if job["early_career"] and job["early_career"] not in ("Yes", "No"):
        raise InvalidJob("early_career must be Yes or No")

    if job["apply_via"]:
        if job["apply_via"] not in APPLY_VIA:
            raise InvalidJob(f"apply_via must be one of {', '.join(APPLY_VIA)}")
    elif record.get("easy_apply") is True:
        job["apply_via"] = "easy-apply"
    elif job["apply_url"]:
        job["apply_via"] = apply_via_from_url(job["apply_url"])
    elif apply_via_from_url(job["url"]) in ("greenhouse", "lever", "ashby", "workday"):
        job["apply_via"] = apply_via_from_url(job["url"])  # the board page is the application
    # otherwise unknown: a LinkedIn/Handshake page link doesn't say how to apply

    ids = [i for i in (job_id_from_url(job["url"]), job_id_from_url(job["apply_url"])) if i]
    ids += [_canonical_id(i.strip().lower()) for i in job["job_id"].split(",") if i.strip()]
    job["job_id"] = ", ".join(dict.fromkeys(ids))
    if not job["date"]:
        job["date"] = date.today().isoformat()
    if not job["status"]:
        job["status"] = DEFAULT_STATUS
    return job


# ── Slugs and job files ──────────────────────────────────────────────

def make_slug(company: str, title: str, location: str = "") -> str:
    """Same rule as prepare_jobs.py and the dashboard, so all three agree."""
    slug = f"{company}_{title}"
    if location:
        slug += f"_{location}"
    slug = re.sub(r"[^\w\-]", "_", slug)
    slug = re.sub(r"_+", "_", slug)
    return slug[:100]


def record_slug(record: dict) -> str:
    """Keep legacy empty-location filenames compatible with the tracker."""
    return make_slug(record["company"], record["title"], record.get("location") or "nan")


def render_job_md(job: dict[str, Any], slug: str) -> str:
    """The jobs/<slug>.md file that the dashboard and the resume skills read."""
    def val(key: str) -> str:
        v = job.get(key, "")
        return ", ".join(v) if isinstance(v, list) else str(v or "")

    lines = [
        f"# Job: {val('title')} at {val('company')}",
        "",
        "## Details",
        f"- **Company:** {val('company')}",
        f"- **Title:** {val('title')}",
        f"- **Location:** {val('location')}",
        f"- **Category:** {val('category')}",
        f"- **Salary:** {val('salary')}",
    ]
    optional = [
        ("Posting URL", "url"),
        ("Apply URL", "apply_url"),
        ("Apply via", "apply_via"),
        ("Found via", "found_via"),
        ("Flags", "flags"),
    ]
    lines += [f"- **{label}:** {val(key)}" for label, key in optional if val(key)]
    lines += [
        "",
        "## Output Path",
        f"`{(Path('output') / slug / 'resume.tex').as_posix()}`",
        "",
        "## Job Description",
        "",
        val("description"),
        "",
    ]
    if val("company_intro"):
        lines += ["## Company Intro", "", val("company_intro"), ""]
    return "\n".join(lines)


# ── Duplicate matching ───────────────────────────────────────────────

_COMPANY_SUFFIX = re.compile(r"\b(the|inc|llc|ltd|corp|corporation|co|company|technologies|technology)\b\.?", re.I)


def company_words(company: str) -> list[str]:
    base = re.sub(r"\(.*?\)", " ", str(company or ""))
    base = _COMPANY_SUFFIX.sub(" ", base)
    return re.findall(r"[a-z0-9]+", base.lower())


def company_key(company: str) -> str:
    return "".join(company_words(company))


def same_company(a: str, b: str) -> bool:
    """Same name, or one is the other's first words: boards say "Domino" or "Noom US" where the
    sheet says "Domino Data Lab" or "Noom". "Zip" and "ZipRecruiter" stay different companies."""
    wa, wb = company_words(a), company_words(b)
    if not wa or not wb:
        return False
    if len(wa) > len(wb):
        wa, wb = wb, wa
    short = "".join(wa)
    return any("".join(wb[:k]) == short for k in range(1, len(wb) + 1))


_PLACE_NOISE = {"remote", "united", "states", "usa", "hub", "hybrid", "onsite", "office", "new", "city",
                "area", "bay", "greater", "metro", "not", "specified", "nan", "anywhere"}


def _place_words(location: str) -> set[str]:
    return {w for w in re.findall(r"[a-z]{3,}", str(location or "").lower()) if w not in _PLACE_NOISE}


def same_place(a: str, b: str) -> bool:
    """The two locations share a city or state name (or one of them doesn't say)."""
    pa, pb = _place_words(a), _place_words(b)
    return not pa or not pb or bool(pa & pb)


def title_key(title: str) -> str:
    return re.sub(r"[^a-z0-9]", "", str(title or "").lower())


# Words that don't make two postings different jobs: years and the many ways of saying "new grad"
_TITLE_NOISE = {"the", "a", "an", "and", "of", "for", "to", "in", "at", "with", "2024", "2025", "2026", "2027", "2028",
                "new", "grad", "grads", "graduate", "graduates", "early", "career", "careers", "entry", "level",
                "university", "college", "program", "start", "full", "time", "fulltime",
                "remote", "hybrid", "onsite", "us", "usa"}
_TITLE_SAME = {"engineering": "engineer", "engineers": "engineer", "i": "1", "ii": "2", "iii": "3",
               "swe": "software engineer", "sde": "software development engineer"}


def title_core(title: str) -> frozenset[str]:
    """The words that identify a role, in any order: "New Grad Software Engineer" and "Software Engineer,
    2027 New Grad" are the same; "Software Engineer II" or "Software Engineer, Backend" are not.
    Words in brackets count: "Software Engineer I (.Net)" is "Software Engineer 1 - .Net"."""
    words = []
    for word in re.findall(r"[a-z0-9+#]+", str(title or "").lower()):
        words += _TITLE_SAME.get(word, word).split()
    return frozenset(w for w in words if w not in _TITLE_NOISE)


def _row_ids(row: dict[str, Any]) -> set[str]:
    ids = {_canonical_id(i.strip().lower()) for i in str(row.get("job_id") or "").split(",")} - {""}
    for key in ("url", "apply_url"):
        jid = job_id_from_url(str(row.get(key) or ""))
        if jid:
            ids.add(jid)
    return ids


def find_duplicate(job: dict[str, Any], rows: list[dict[str, Any]]) -> tuple[int, str] | None:
    """(row index, reason) of the first stored row that is the same job."""
    ids = _row_ids(job)
    if ids:
        for i, row in enumerate(rows):
            if ids & _row_ids(row):
                return i, "Same job link"
    ck, tk = company_key(job["company"]), title_key(job["title"])
    if ck and tk:
        for i, row in enumerate(rows):
            if company_key(row.get("company", "")) == ck and title_key(row.get("title", "")) == tk:
                return i, "Same company and title"
        core = title_core(job["title"])
        if len(core) >= 2:  # the same role worded differently (word order, year, "New Grad" vs "New Graduate")
            for i, row in enumerate(rows):
                if title_core(row.get("title", "")) != core:
                    continue
                if company_key(row.get("company", "")) == ck:
                    return i, "Same company and title, worded differently"
                # "Domino" vs "Domino Data Lab": the place has to agree too, so "Apple" and "Apple Bank" don't merge
                if same_company(job["company"], row.get("company", "")) and same_place(job.get("location"), row.get("location")):
                    return i, "Same company (named differently) and title"
    return None


# ── Store: jobs.xlsx + jobs/<slug>.md + dashboard_statuses.json ───────

class Store:
    def __init__(self, root: Path = REPO_ROOT):
        self.root = Path(root)
        self.xlsx = self.root / "jobs.xlsx"
        self.jobs_dir = self.root / "jobs"
        self.statuses = self.root / "dashboard_statuses.json"
        self.lock_file = self.root / ".jobs.lock"
        self.backup_dir = self.root / "archive"

    @contextlib.contextmanager
    def locked(self) -> Iterator[None]:
        """One writer at a time across the dashboard, schedules and sweeps."""
        with open(self.lock_file, "a+") as fh:
            fcntl.flock(fh, fcntl.LOCK_EX)
            try:
                yield
            finally:
                fcntl.flock(fh, fcntl.LOCK_UN)

    def _load(self):
        from openpyxl import Workbook, load_workbook

        if self.xlsx.exists():
            wb = load_workbook(self.xlsx)
        else:
            wb = Workbook()
            wb.active.title = "Sheet1"
            wb.active.append(list(COLUMNS.values()))
        ws = wb.active
        headers = [str(c.value) if c.value is not None else "" for c in ws[1]]
        return wb, ws, headers

    def _ensure_columns(self, ws, headers: list[str]) -> list[str]:
        missing = [col for col in COLUMNS.values() if col not in headers]
        if not missing:
            return headers
        if self.xlsx.exists():
            self.backup_dir.mkdir(exist_ok=True)
            stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
            shutil.copy2(self.xlsx, self.backup_dir / f"jobs.backup-{stamp}.xlsx")
        for col in missing:
            ws.cell(row=1, column=len(headers) + 1, value=col)
            headers.append(col)
        return headers

    @staticmethod
    def _row_to_record(values: tuple, headers: list[str]) -> dict[str, Any]:
        by_col = {h: values[i] if i < len(values) else None for i, h in enumerate(headers)}
        rec: dict[str, Any] = {}
        for key, col in COLUMNS.items():
            v = by_col.get(col)
            v = "" if v is None else str(v).strip()
            rec[key] = [t for t in re.split(r"\s*,\s*", v) if t] if key in LIST_FIELDS else v
        return rec

    def _save_atomic(self, wb) -> None:
        fd, tmp = tempfile.mkstemp(dir=self.root, prefix=".jobs-", suffix=".xlsx")
        os.close(fd)
        try:
            wb.save(tmp)
            os.replace(tmp, self.xlsx)
        finally:
            if os.path.exists(tmp):
                os.unlink(tmp)

    def _write_text_atomic(self, path: Path, text: str) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        fd, tmp = tempfile.mkstemp(dir=path.parent, prefix=".tmp-")
        with os.fdopen(fd, "w", encoding="utf-8") as fh:
            fh.write(text)
        os.replace(tmp, path)

    def _set_status(self, slug: str, status: str) -> None:
        statuses: dict[str, str] = {}
        if self.statuses.exists():
            try:
                statuses = json.loads(self.statuses.read_text(encoding="utf-8")) or {}
            except (OSError, json.JSONDecodeError):
                statuses = {}
        statuses[slug] = status
        self._write_text_atomic(self.statuses, json.dumps(statuses, indent=2, sort_keys=True))

    def rows(self) -> list[dict[str, Any]]:
        if not self.xlsx.exists():
            return []
        _, ws, headers = self._load()
        return [self._row_to_record(v, headers) for v in ws.iter_rows(min_row=2, values_only=True) if any(v)]

    def add(self, record: dict[str, Any], *, force: bool = False, dry_run: bool = False,
            capture_undo: bool = False) -> dict[str, Any]:
        try:
            job = normalize(record)
        except InvalidJob as exc:
            return {"result": "rejected", "reason": str(exc)}

        slug = make_slug(job["company"], job["title"], job["location"])
        with self.locked():
            wb, ws, headers = self._load()
            rows = [self._row_to_record(v, headers) for v in ws.iter_rows(min_row=2, values_only=True)]
            dup = find_duplicate(job, rows)

            if dup and not force:
                idx, reason = dup
                existing = rows[idx]
                before = json.loads(json.dumps(existing))
                existing_slug = record_slug(existing)
                md = self.jobs_dir / f"{existing_slug}.md"
                before_md = md.read_text(encoding="utf-8") if md.exists() else None
                changes = self._merge(existing, job)
                result = {
                    "result": "merged" if changes else "duplicate",
                    "reason": reason,
                    "slug": make_slug(existing["company"], existing["title"], existing["location"] or "nan"),
                    "changed": changes,
                    "existing": {COLUMNS[k]: (", ".join(v) if isinstance(v, list) else v) for k, v in existing.items()},
                }
                if dry_run:
                    result["dry_run"] = True
                elif changes:
                    headers = self._ensure_columns(ws, headers)
                    self._write_row(ws, headers, idx + 2, existing, keys=changes)
                    self._save_atomic(wb)
                    existing_slug = result["slug"]
                    md = self.jobs_dir / f"{existing_slug}.md"
                    if md.exists():
                        self._write_text_atomic(md, render_job_md(existing, existing_slug))
                if capture_undo and not dry_run:
                    result["_undo"] = {
                        "kind": result["result"], "slug": result["slug"],
                        "before": before, "after": self._row_to_record(tuple(c.value for c in ws[idx + 2]), headers),
                        "before_md": before_md,
                        "after_md": md.read_text(encoding="utf-8") if md.exists() else None,
                    }
                return result

            if any(make_slug(r["company"], r["title"], r["location"] or "nan") == slug for r in rows):
                return {"result": "rejected", "slug": slug,
                        "reason": "a job with the same company, title and location already exists"}
            if dry_run:
                return {"result": "added", "slug": slug, "dry_run": True, "job": job}

            previous_md = self.jobs_dir / f"{slug}.md"
            before_md = previous_md.read_text(encoding="utf-8") if previous_md.exists() else None
            before_status = None
            if capture_undo and self.statuses.exists():
                before_status = json.loads(self.statuses.read_text(encoding="utf-8")).get(slug)
            headers = self._ensure_columns(ws, headers)
            self._write_row(ws, headers, ws.max_row + 1, job)
            self._save_atomic(wb)
            self._write_text_atomic(self.jobs_dir / f"{slug}.md", render_job_md(job, slug))
            self._set_status(slug, job["status"])
            result = {"result": "added", "slug": slug, "job_id": job["job_id"], "apply_via": job["apply_via"]}
            if capture_undo:
                result["_undo"] = {
                    "kind": "added", "slug": slug, "before": None, "before_md": before_md,
                    "before_status": before_status,
                    "after": self._row_to_record(tuple(c.value for c in ws[ws.max_row]), headers),
                    "after_md": (self.jobs_dir / f"{slug}.md").read_text(encoding="utf-8"),
                }
            return result

    def _find_row(self, ws, headers, slug):
        for idx, values in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            row = self._row_to_record(values, headers)
            if any(values) and record_slug(row) == slug:
                return idx, row
        raise KeyError("Job not found")

    def set_status(self, slug: str, status: str) -> None:
        with self.locked():
            wb, ws, headers = self._load()
            idx, _ = self._find_row(ws, headers, slug)
            headers = self._ensure_columns(ws, headers)
            self._write_row(ws, headers, idx, {"status": status}, keys=["status"])
            self._save_atomic(wb)
            self._set_status(slug, status)

    def _remove_status(self, slug: str) -> None:
        if self.statuses.exists():
            statuses = json.loads(self.statuses.read_text(encoding="utf-8"))
            statuses.pop(slug, None)
            self._write_text_atomic(self.statuses, json.dumps(statuses, indent=2, sort_keys=True))

    def delete(self, slug: str) -> dict:
        """Remove a tracker row and retain its files in a local recovery folder."""
        with self.locked():
            wb, ws, headers = self._load()
            idx, row = self._find_row(ws, headers, slug)
            recovery = self.root / "data" / "trash" / f"{datetime.now():%Y%m%d-%H%M%S}-{uuid.uuid4().hex[:8]}"
            recovery.mkdir(parents=True)
            self._write_text_atomic(recovery / "record.json", json.dumps(row, indent=2))
            shutil.copy2(self.xlsx, recovery / "jobs.xlsx")
            ws.delete_rows(idx)
            self._save_atomic(wb)
            removed = {"xlsx_row": True, "job_md": False, "output_dir": False}
            for key, path in (("job_md", self.jobs_dir / f"{slug}.md"), ("output_dir", self.root / "output" / slug)):
                if path.exists():
                    shutil.move(str(path), str(recovery / path.name))
                    removed[key] = True
            self._remove_status(slug)
            return removed

    def undo_keep(self, snapshot: dict) -> None:
        """Undo only this Keep's changes. Never remove a pre-existing duplicate.

        Refuse if someone has edited the job since Keep, or generated artifacts.
        This prevents a stale undo from erasing later work.
        """
        if snapshot["kind"] == "duplicate":
            return
        slug = snapshot["slug"]
        with self.locked():
            wb, ws, headers = self._load()
            idx, current = self._find_row(ws, headers, slug)
            md = self.jobs_dir / f"{slug}.md"
            current_md = md.read_text(encoding="utf-8") if md.exists() else None
            if current != snapshot["after"] or current_md != snapshot["after_md"]:
                raise ValueError("This job has changed since Keep. It was left in My jobs.")
            if snapshot["kind"] == "added":
                out = self.root / "output" / slug
                if out.exists() and any(out.iterdir()):
                    raise ValueError("This job now has files. It was left in My jobs.")
                if self.statuses.exists():
                    status = json.loads(self.statuses.read_text(encoding="utf-8")).get(slug)
                    if status != snapshot["after"]["status"]:
                        raise ValueError("This job's status has changed. It was left in My jobs.")
                ws.delete_rows(idx)
            else:
                # openpyxl's cell(value=None) doesn't clear a cell, so assign directly.
                for key, col in COLUMNS.items():
                    if col in headers:
                        value = snapshot["before"].get(key, "")
                        ws.cell(idx, headers.index(col) + 1).value = ", ".join(value) if isinstance(value, list) else value
            self._save_atomic(wb)
            if snapshot["before_md"] is None:
                md.unlink(missing_ok=True)
            else:
                self._write_text_atomic(md, snapshot["before_md"])
            if snapshot["kind"] == "added":
                if snapshot.get("before_status") is not None:
                    self._set_status(slug, snapshot["before_status"])
                else:
                    self._remove_status(slug)

    @staticmethod
    def _merge(existing: dict[str, Any], new: dict[str, Any]) -> list[str]:
        """Add new sources and fill empty fields. Never overwrites what's there."""
        changes: list[str] = []
        for key in LIST_FIELDS:
            added = [v for v in new[key] if v not in existing[key]]
            if added:
                existing[key] = existing[key] + added
                changes.append(key)
        new_ids = [i for i in new["job_id"].split(", ") if i and i not in existing["job_id"].split(", ")]
        if new_ids:
            existing["job_id"] = ", ".join([i for i in existing["job_id"].split(", ") if i] + new_ids)
            changes.append("job_id")
        # Location participates in the filename; merging must not silently change the slug.
        for key in COLUMNS:
            if key in LIST_FIELDS or key in ("status", "date", "job_id", "location"):
                continue
            if not existing.get(key) and new.get(key):
                existing[key] = new[key]
                changes.append(key)
        # A real apply link beats "other"/unknown for how to apply
        if existing.get("apply_via") == "other" and new.get("apply_via") not in ("", "other") and "apply_via" not in changes:
            existing["apply_via"] = new["apply_via"]
            changes.append("apply_via")
        return changes

    @staticmethod
    def _write_row(ws, headers: list[str], row_idx: int, job: dict[str, Any], keys: list[str] | None = None) -> None:
        """Write a job's cells; with keys, only those cells (merges leave the rest untouched)."""
        col_for_key = {col: key for key, col in COLUMNS.items()}
        for i, header in enumerate(headers, start=1):
            key = col_for_key.get(header)
            if key is None or (keys is not None and key not in keys):
                continue
            value = job.get(key, "")
            if isinstance(value, list):
                value = ", ".join(value)
            ws.cell(row=row_idx, column=i, value=value if value != "" else None)


def record_from_columns(fields: dict[str, Any]) -> dict[str, Any]:
    """Convert a jobs.xlsx-column dict (the dashboard form) to a record."""
    key_for_col = {col: key for key, col in COLUMNS.items()}
    return {key_for_col[col]: value for col, value in fields.items() if col in key_for_col}


def columns_from_record(record: dict[str, Any]) -> dict[str, str]:
    """Convert a record to jobs.xlsx-column names (for the dashboard form)."""
    out = {}
    for key, value in record.items():
        if key in COLUMNS:
            out[COLUMNS[key]] = ", ".join(value) if isinstance(value, list) else ("" if value is None else str(value))
    return out


def add_job(record: dict[str, Any], *, root: Path = REPO_ROOT, force: bool = False, dry_run: bool = False) -> dict[str, Any]:
    return Store(root).add(record, force=force, dry_run=dry_run)


# ── CLI ──────────────────────────────────────────────────────────────

def _read_records(source: str) -> list[Any]:
    text = sys.stdin.read() if source == "-" else Path(source).read_text(encoding="utf-8")
    text = text.strip()
    if not text:
        return []
    if text[0] in "[{":
        try:
            data = json.loads(text)
            return data if isinstance(data, list) else [data]
        except json.JSONDecodeError:
            pass
    return [json.loads(line) for line in text.splitlines() if line.strip()]


def main() -> int:
    parser = argparse.ArgumentParser(description="Add job records to Loom.")
    parser.add_argument("source", help="JSON/JSONL file, or - for stdin")
    parser.add_argument("--force", action="store_true", help="add even if it matches an existing job")
    parser.add_argument("--dry-run", action="store_true", help="validate and check duplicates without writing")
    parser.add_argument("--root", type=Path, default=REPO_ROOT, help=argparse.SUPPRESS)
    args = parser.parse_args()

    try:
        records = _read_records(args.source)
    except (OSError, json.JSONDecodeError) as exc:
        print(json.dumps({"result": "rejected", "reason": f"couldn't read input: {exc}"}))
        return 1

    store = Store(args.root)
    failed = 0
    for record in records:
        result = store.add(record, force=args.force, dry_run=args.dry_run)
        result.pop("existing", None)
        result.pop("job", None)
        failed += result["result"] == "rejected"
        print(json.dumps(result, ensure_ascii=False))
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
